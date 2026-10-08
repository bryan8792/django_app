from datetime import date
from decimal import Decimal, InvalidOperation
import json
import ssl
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from urllib.parse import urlparse
from django.utils import timezone
import re
import unicodedata
from openpyxl import load_workbook
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.http import HttpResponseForbidden, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
import os
import secrets
class StaffRequiredMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_staff

    def handle_no_permission(self):
        if self.request.user.is_authenticated:
            return HttpResponseForbidden('No tienes permisos para esta sección.')
        return super().handle_no_permission()
from django.db import transaction
from django.db.models import Q, Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse_lazy
from django.views import View
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView
from django.contrib.auth import get_user_model
from .forms import CampaignForm, ExcelImportForm, InvoiceForm, StaffUserCreationForm, StaffUserUpdateForm

User = get_user_model()
from .models import Campaign, ImportBatch, Invoice, MessageLog, Subscriber

class DashboardView(LoginRequiredMixin, TemplateView):
    template_name = 'billing/dashboard.html'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(
            subscribers=Subscriber.objects.filter(active=True).count(),
            debtors=Subscriber.objects.filter(invoices__status='pending', invoices__amount__gt=0).distinct().count(),
            debt=Invoice.objects.filter(status='pending').aggregate(total=Sum('amount'))['total'] or Decimal('0'),
            recent=Invoice.objects.select_related('subscriber')[:8],
        )
        return context

class SubscriberListView(LoginRequiredMixin, ListView):
    model = Subscriber
    template_name = 'billing/subscriber_list.html'
    paginate_by = 25
    context_object_name = 'subscribers'
    def get_queryset(self):
        query = self.request.GET.get('q', '').strip()
        subscribers = Subscriber.objects.order_by('full_name')
        if not query:
            return subscribers
        return subscribers.filter(Q(full_name__icontains=query) | Q(account_number__icontains=query) | Q(phone__icontains=query))

class SubscriberDetailView(LoginRequiredMixin, DetailView):
    model = Subscriber
    template_name = 'billing/subscriber_detail.html'

class InvoiceUpdateView(StaffRequiredMixin, LoginRequiredMixin, UpdateView):
    model = Invoice
    form_class = InvoiceForm
    template_name = 'billing/form.html'
    success_url = reverse_lazy('subscribers')

class InvoicePreviewView(LoginRequiredMixin, DetailView):
    model = Invoice
    template_name = 'billing/invoice_preview.html'
    context_object_name = 'invoice'

    def get_queryset(self):
        return Invoice.objects.select_related('subscriber')

def normalize_column(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]', '', text)


def normalize_phone(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
        if 'E+' in text.upper():
            try:
                text = str(int(Decimal(text)))
            except InvalidOperation:
                pass
    digits = re.sub(r'\D', '', text)
    if digits.startswith('00'):
        digits = digits[2:]
    if digits.startswith('0') and len(digits) == 10:
        digits = '593' + digits[1:]
    elif digits.startswith('9') and len(digits) == 9:
        digits = '593' + digits
    return digits


class ImportExcelView(StaffRequiredMixin, LoginRequiredMixin, View):
    template_name = 'billing/import.html'
    def get(self, request):
        latest_batch = ImportBatch.objects.order_by('-created_at').first()
        return render(request, self.template_name, {'form': ExcelImportForm(), 'latest_batch': latest_batch})
    @transaction.atomic
    def post(self, request):
        form = ExcelImportForm(request.POST, request.FILES)
        if not form.is_valid():
            return render(request, self.template_name, {'form': form})
        uploaded = form.cleaned_data['file']
        batch = ImportBatch.objects.create(file=uploaded, created_by=request.user)
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        header_row = next(sheet.iter_rows())
        headers = [normalize_column(cell.value) for cell in header_row]
        aliases = {
            'CUENTA': ['CUENTA', 'NUMERODECUENTA', 'NROCUENTA', 'NUMCUENTA'],
            'NOMBRE': ['NOMBRE', 'ABONADO', 'PROPIETARIO', 'USUARIO', 'NOMBREABONADO', 'ABONADOCLIE'],
            'TELEFONO': ['TELEFONO', 'WHATSAPP', 'CELULAR', 'MOVIL', 'TELEFONOCELULAR', 'NUMEROTELEFONO', 'DECELULAR', 'NDELCELULAR', 'NUMEROCELULAR'],
            'EMAIL': ['CORREO', 'EMAIL', 'CORREOELECTRONICO'],
            'DEUDA': ['TOTAL', 'DEUDA', 'VALOR', 'TOTALDEUDA', 'SALDO'],
            'CEDULA': ['CEDULA', 'IDENTIFICACION', 'DOCUMENTO'],
            'DIRECCION': ['DIRECCION', 'DOMICILIO', 'DIRECCIONUBI'],
            'CATEGORIA': ['CATEGORIA', 'CATEGORI', 'TIPOCLIENTE'],
            'SECUENCIA': ['SECUENCIA', 'SECUENCI'],
            'CODIGO': ['ABONADOCLIENTE', 'ABONADOCLIE', 'CODIGOABONADO'],
            'MEDIDOR': ['NROMEDIDOR', 'NROMED'],
            'LECTURA_ANTERIOR': ['LECTURAANTERIOR', 'LECTURAANTE'],
            'LECTURA_ACTUAL': ['LECTURAACTUAL', 'LECTURAACT'],
            'CARTAS': ['TOTALCARTAS', 'TOTALCART'],
        }
        indexes = {key: next((headers.index(normalize_column(alias)) for alias in names if normalize_column(alias) in headers), None) for key, names in aliases.items()}
        imported = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            def value(key):
                index = indexes[key]
                return row[index] if index is not None and index < len(row) else None
            account = str(value('CUENTA') or '').strip()
            name = str(value('NOMBRE') or '').strip()
            if not account or not name:
                continue
            subscriber, _ = Subscriber.objects.update_or_create(
                account_number=account,
                defaults={
                    'import_batch': batch,
                    'full_name': name,
                    'phone': normalize_phone(value('TELEFONO')),
                    'email': str(value('EMAIL') or '').strip(),
                    'national_id': str(value('CEDULA') or '').strip(),
                    'address': str(value('DIRECCION') or '').strip(),
                    'category': str(value('CATEGORIA') or '').strip(),
                    'sequence': str(value('SECUENCIA') or '').strip(),
                    'client_code': str(value('CODIGO') or '').strip(),
                    'meter_number': str(value('MEDIDOR') or '').strip(),
                    'previous_reading': Decimal(str(value('LECTURA_ANTERIOR') or 0).replace(',', '.')),
                    'current_reading': Decimal(str(value('LECTURA_ACTUAL') or 0).replace(',', '.')),
                    'total_letters': Decimal(str(value('CARTAS') or 0).replace(',', '.')),
                },
            )
            try:
                amount = Decimal(str(value('DEUDA') or 0).replace(',', '.'))
            except (InvalidOperation, ValueError):
                amount = Decimal('0')
            Invoice.objects.update_or_create(subscriber=subscriber, period=date.today().replace(day=1), defaults={'amount': amount})
            imported += 1
        batch.rows_imported = imported
        batch.save(update_fields=['rows_imported'])
        messages.success(request, f'Se importaron {imported} abonados e información de factura.')
        return redirect('subscribers')

class WhatsAppQrView(StaffRequiredMixin, LoginRequiredMixin, View):
    def get(self, request):
        api_url = os.getenv('WHATSAPP_QR_API', 'https://api-msg.alesof.xyz/whatsapp/qr')
        try:
            with urlopen(Request(api_url, headers={'Accept': 'image/png, image/jpeg, application/json'}), timeout=20) as response:
                content = response.read()
                content_type = response.headers.get('Content-Type', 'image/png')
            return HttpResponse(content, content_type=content_type)
        except (HTTPError, URLError, TimeoutError) as exc:
            return JsonResponse({'ok': False, 'error': str(exc)}, status=502)

class CampaignCreateView(StaffRequiredMixin, LoginRequiredMixin, CreateView):
    model = Campaign
    form_class = CampaignForm
    template_name = 'billing/campaign_form.html'
    success_url = reverse_lazy('campaigns')
    def form_valid(self, form):
        recipients = form.cleaned_data.pop('recipients')
        if form.cleaned_data.get('send_mode') == 'now':
            form.instance.scheduled_for = timezone.now()
        form.instance.created_by = self.request.user
        response = super().form_valid(form)
        MessageLog.objects.bulk_create([MessageLog(campaign=self.object, subscriber=subscriber) for subscriber in recipients])
        if form.cleaned_data.get('send_mode') == 'now':
            return redirect('campaign_queue', pk=self.object.pk)
        return response

class CampaignListView(StaffRequiredMixin, LoginRequiredMixin, ListView):
    model = Campaign
    template_name = 'billing/campaign_list.html'
    context_object_name = 'campaigns'

class UserListView(StaffRequiredMixin, LoginRequiredMixin, ListView):
    model = User
    template_name = 'billing/user_list.html'
    context_object_name = 'users'
    def get_queryset(self):
        query = self.request.GET.get('q', '').strip()
        users = User.objects.order_by('username')
        if query:
            users = users.filter(Q(username__icontains=query) | Q(first_name__icontains=query) | Q(last_name__icontains=query) | Q(email__icontains=query))
        return users

class UserCreateView(StaffRequiredMixin, LoginRequiredMixin, CreateView):
    model = User
    form_class = StaffUserCreationForm
    template_name = 'billing/user_form.html'
    success_url = reverse_lazy('users')

class UserUpdateView(StaffRequiredMixin, LoginRequiredMixin, UpdateView):
    model = User
    form_class = StaffUserUpdateForm
    template_name = 'billing/user_form.html'
    success_url = reverse_lazy('users')

class UserToggleView(StaffRequiredMixin, LoginRequiredMixin, View):
    def post(self, request, pk):
        user = get_object_or_404(User, pk=pk)
        if user != request.user:
            user.is_active = not user.is_active
            user.save(update_fields=['is_active'])
        return redirect('users')

class CampaignQueueView(StaffRequiredMixin, LoginRequiredMixin, View):
    def get(self, request, pk):
        return self._send(request, pk)

    def post(self, request, pk):
        return self._send(request, pk)

    def _send(self, request, pk):
        campaign = get_object_or_404(Campaign, pk=pk)
        recipients = campaign.messages.select_related('subscriber').filter(status__in=['queued', 'failed'])
        sent = 0
        failed = 0
        for log in recipients:
            phone = re.sub(r'\D', '', log.subscriber.phone or '')
            country_code = os.getenv('WA_COUNTRY_CODE', '593').strip()
            if phone.startswith('00'):
                phone = phone[2:]
            elif country_code and len(phone) == 9:
                phone = country_code + phone
            invoice = log.subscriber.invoices.filter(period=campaign.period).first() or log.subscriber.invoices.order_by('-period').first()
            amount = invoice.amount if invoice else Decimal('0.00')
            spanish_months = ('enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre')
            period_name = spanish_months[campaign.period.month - 1] if campaign.period else ''
            text = campaign.message_template.replace('{nombre}', log.subscriber.full_name).replace('{cuenta}', log.subscriber.account_number).replace('{total}', f'{amount:.2f}').replace('{periodo}', period_name)
            payload = json.dumps({'to': phone, 'message': text}).encode('utf-8')
            try:
                api_url = os.getenv('WHATSAPP_SEND_API', 'https://api-msg.alesof.xyz/whatsapp/send')
                api_request = Request(api_url, data=payload, headers={'Content-Type': 'application/json', 'Accept': 'application/json'}, method='POST')
                # Esta API usa el mismo certificado autofirmado que Postman acepta.
                # Se limita esta excepción a la conexión configurada de WhatsApp.
                insecure_context = ssl._create_unverified_context()
                with urlopen(api_request, timeout=30, context=insecure_context) as api_response:
                    result = json.loads(api_response.read().decode('utf-8'))
                if not result.get('ok'):
                    raise ValueError(result.get('error') or 'La API no confirmó el envío.')
                log.status = 'sent'
                log.sent_at = timezone.now()
                log.error = ''
                sent += 1
            except HTTPError as exc:
                detail = exc.read().decode('utf-8', errors='replace')
                log.status = 'failed'
                log.error = f'WhatsApp API {exc.code}: {detail[:500]}'
                failed += 1
            except (URLError, KeyError, TimeoutError, ValueError) as exc:
                log.status = 'failed'
                log.error = str(exc)
                failed += 1
            except Exception as exc:
                log.status = 'failed'
                log.error = f'{type(exc).__name__}: {exc}'
                failed += 1
            log.save(update_fields=['status', 'sent_at', 'error'])
        campaign.status = 'sent' if sent > 0 and failed == 0 else ('partial' if sent > 0 else 'failed')
        campaign.save(update_fields=['status'])
        messages.success(request, f'{sent} mensajes enviados y {failed} con error.')
        return redirect('campaigns')
